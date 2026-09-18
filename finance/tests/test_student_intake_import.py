from io import BytesIO

from django.contrib.auth.models import Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook

from finance.models import StudentIntakeImportLog, StudentIntakeTrend
from finance.services.student_registration_import import (
    ImportValidationError,
    import_rows,
    preview_upload,
)

HEADERS = [
    'Kode Prodi', 'Fakultas / Kampus', 'Program Studi', 'Tahun',
    'Tarif', 'Kuota', 'Registrasi',
]


def workbook_file(rows, *, headers=None, sheet='Template Upload', filename='upload.xlsx'):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(headers or HEADERS)
    for row in rows:
        worksheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return SimpleUploadedFile(
        filename, stream.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


class StudentIntakeImportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('importer', password='secret')

    def test_valid_preview_classifies_new_and_confirm_upserts(self):
        upload = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, 11000000, 80, 38],
        ])
        rows, summary = preview_upload(upload)
        self.assertEqual(summary, {'total': 1, 'valid': 1, 'new': 1, 'update': 0, 'error': 0})
        self.assertEqual(rows[0]['status'], 'NEW')
        result = import_rows(rows, user=self.user, source_file_name=upload.name)
        self.assertEqual(result, {'new': 1, 'updated': 0, 'total': 1})
        row = StudentIntakeTrend.objects.get(study_program_code='5101', year=2026)
        self.assertEqual(row.quota, 80)
        self.assertEqual(row.registration, 38)
        self.assertEqual(str(row.tariff), '11000000.00')

        upload2 = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, 12000000, 90, 70],
        ], filename='second.xlsx')
        rows2, summary2 = preview_upload(upload2)
        self.assertEqual(summary2['update'], 1)
        self.assertEqual(rows2[0]['status'], 'UPDATE')
        result2 = import_rows(rows2, user=self.user, source_file_name=upload2.name)
        self.assertEqual(result2, {'new': 0, 'updated': 1, 'total': 1})
        row.refresh_from_db()
        self.assertEqual(row.quota, 90)
        self.assertEqual(row.registration, 70)
        self.assertEqual(str(row.tariff), '12000000.00')

    def test_duplicate_key_inside_excel_is_error_and_cannot_import(self):
        upload = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, 11000000, 80, 38],
            ['5101', 'FTE', 'Teknik Elektro', 2026, 11000000, 81, 39],
        ])
        rows, summary = preview_upload(upload)
        self.assertEqual(summary['error'], 2)
        self.assertTrue(all(r['status'] == 'ERROR' for r in rows))
        with self.assertRaisesMessage(ImportValidationError, 'ERROR'):
            import_rows(rows, user=self.user, source_file_name=upload.name)
        self.assertEqual(StudentIntakeTrend.objects.count(), 0)

    def test_missing_headers_are_rejected(self):
        upload = workbook_file(
            [['5101', 'FTE', 'Teknik Elektro', 2026, 11000000, 80]],
            headers=HEADERS[:-1],
        )
        with self.assertRaisesMessage(ImportValidationError, 'Registrasi'):
            preview_upload(upload)

    def test_empty_tariff_stays_null_and_counts_become_zero(self):
        upload = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, None, None, None],
        ])
        rows, summary = preview_upload(upload)
        self.assertEqual(summary['valid'], 1)
        self.assertIsNone(rows[0]['tariff'])
        self.assertEqual(rows[0]['quota'], 0)
        self.assertEqual(rows[0]['registration'], 0)

    def test_formula_cells_are_rejected(self):
        upload = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, '=11000000', 80, 38],
        ])
        rows, summary = preview_upload(upload)
        self.assertEqual(summary['error'], 1)
        self.assertIn('formula', rows[0]['validation'].lower())

    def test_wrong_extension_and_malformed_xlsx_are_rejected(self):
        bad_ext = SimpleUploadedFile('upload.xls', b'not excel')
        with self.assertRaisesMessage(ImportValidationError, '.xlsx'):
            preview_upload(bad_ext)
        malformed = SimpleUploadedFile('upload.xlsx', b'not a zip')
        with self.assertRaisesMessage(ImportValidationError, 'valid'):
            preview_upload(malformed)

    def test_missing_database_rows_are_preserved_by_upsert(self):
        StudentIntakeTrend.objects.create(
            study_program_code='OLD', faculty_name='FTE',
            study_program_name='Old Program', year=2019,
            quota=10, registration=5, tariff=100,
            source_file_name='old.xlsx', uploaded_by=self.user,
        )
        upload = workbook_file([
            ['NEW', 'FIF', 'New Program', 2026, 1000, 10, 9],
        ])
        rows, _ = preview_upload(upload)
        import_rows(rows, user=self.user, source_file_name=upload.name)
        self.assertTrue(StudentIntakeTrend.objects.filter(study_program_code='OLD').exists())
        self.assertTrue(StudentIntakeTrend.objects.filter(study_program_code='NEW').exists())

    def test_preview_and_confirm_view_are_permission_gated(self):
        viewer = self.client
        upload = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, 11000000, 80, 38],
        ])
        viewer.force_login(self.user)
        response = viewer.post('/dashboard/registrasi-mahasiswa/upload/preview/', {'file': upload})
        self.assertEqual(response.status_code, 403)

        permission = Permission.objects.get(
            content_type__app_label='finance', codename='add_studentintaketrend'
        )
        self.user.user_permissions.add(permission)
        upload = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, 11000000, 80, 38],
        ])
        response = viewer.post('/dashboard/registrasi-mahasiswa/upload/preview/', {'file': upload})
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload['ok'])
        self.assertEqual(payload['summary']['new'], 1)
        self.assertFalse(payload['summary']['error'])
        self.assertEqual(StudentIntakeImportLog.objects.count(), 1)

        response = viewer.post(
            '/dashboard/registrasi-mahasiswa/upload/confirm/',
            {'token': payload['token']},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['new'], 1)
        self.assertEqual(StudentIntakeTrend.objects.count(), 1)
        self.assertEqual(StudentIntakeImportLog.objects.get().status, 'SUCCESS')

    def test_invalid_preview_never_writes_rows(self):
        permission = Permission.objects.get(
            content_type__app_label='finance', codename='add_studentintaketrend'
        )
        self.user.user_permissions.add(permission)
        self.client.force_login(self.user)
        upload = workbook_file([
            ['5101', 'FTE', 'Teknik Elektro', 2026, -1, 80, 38],
        ])
        response = self.client.post(
            '/dashboard/registrasi-mahasiswa/upload/preview/', {'file': upload}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['summary']['error'], 1)
        self.assertEqual(StudentIntakeTrend.objects.count(), 0)

    def test_template_download_has_canonical_sheets_and_headers(self):
        permission = Permission.objects.get(
            content_type__app_label='finance', codename='add_studentintaketrend'
        )
        self.user.user_permissions.add(permission)
        self.client.force_login(self.user)
        response = self.client.get('/dashboard/registrasi-mahasiswa/template/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('Template_Upload_Registrasi_Kuota_BPP.xlsx', response['Content-Disposition'])
        import openpyxl
        workbook = openpyxl.load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ['Template Upload', 'Petunjuk'])
        self.assertEqual(
            list(next(workbook['Template Upload'].iter_rows(values_only=True))),
            HEADERS,
        )
        workbook.close()

    def test_viewer_can_view_but_cannot_see_upload_button(self):
        self.client.force_login(self.user)
        response = self.client.get('/dashboard/registrasi-mahasiswa/')
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'Upload Data Excel')
        self.assertNotContains(response, 'Download Template Excel')

    def test_upload_size_is_rejected(self):
        permission = Permission.objects.get(
            content_type__app_label='finance', codename='add_studentintaketrend'
        )
        self.user.user_permissions.add(permission)
        self.client.force_login(self.user)
        big = SimpleUploadedFile('big.xlsx', b'x' * (10 * 1024 * 1024 + 1))
        response = self.client.post(
            '/dashboard/registrasi-mahasiswa/upload/preview/', {'file': big}
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('10 MB', response.json()['message'])
