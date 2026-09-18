"""Registrasi Mahasiswa analysis and Excel upsert endpoints."""
import json
import time
from io import BytesIO
from uuid import uuid4

from django.db import OperationalError, ProgrammingError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from dashboard.views import _assets_head, _fonts_head

from .models import StudentIntakeImportLog
from .permissions import can_upload_student_intake, require_student_intake_upload
from .services import student_registration_import as importer
from .services import student_registration_service as srs

ACTIVE_TAB = 'registrasi_mahasiswa'
PREVIEW_SESSION_KEY = 'student_intake_preview'


def _scope(request):
    """Read the faculty/program filter from the query string.

    The selection is applied exactly as given. If it matches nothing the page
    shows its empty state, which is honest: silently widening an unknown filter
    to the full dataset would let an operator read unfiltered totals while
    believing they were looking at one faculty.
    """
    return (
        (request.GET.get('faculty') or '').strip(),
        (request.GET.get('study_program') or '').strip(),
    )


def _recent_imports():
    try:
        return StudentIntakeImportLog.objects.select_related('uploaded_by')[:10]
    except (OperationalError, ProgrammingError):
        return []


def _context(request, faculty, study_program, trend):
    chart_label = srs.filter_labels(faculty, study_program)
    return {
        'assets_head': _assets_head(),
        'fonts_head': _fonts_head(),
        'active': 'dashboard',
        'active_tab': ACTIVE_TAB,
        'faculties': srs.get_faculties(),
        'study_programs': srs.get_study_programs(faculty or None),
        'selected_faculty': faculty,
        'selected_study_program': study_program,
        'trend': trend,
        'chart_label': chart_label,
        'can_upload': can_upload_student_intake(request.user),
        'template_url': reverse('registration:template'),
        'upload_preview_url': reverse('registration:upload-preview'),
        'upload_confirm_url': reverse('registration:upload-confirm'),
        'import_logs': _recent_imports(),
        # Same payload the JSON endpoint returns, inlined so the first paint
        # needs no extra round trip (and works before the JS runs).
        'trend_json': json.dumps({
            'years': trend['years'],
            'series': trend['series'],
            'summary': trend['summary'],
            'has_data': trend['has_data'],
            'table': trend['table'],
            'chart_label': chart_label,
        }),
    }


def registrasi_mahasiswa(request):
    faculty, study_program = _scope(request)
    trend = srs.get_registration_trend(faculty or None, study_program or None)
    return render(request, 'finance/registrasi_mahasiswa.html',
                  _context(request, faculty, study_program, trend))


def registrasi_mahasiswa_data(request):
    """Aggregated series for the current filter (JSON)."""
    faculty, study_program = _scope(request)
    return JsonResponse(srs.get_registration_trend(faculty or None, study_program or None))


def registrasi_mahasiswa_programs(request):
    """Study programs within a faculty, for the dependent dropdown (JSON)."""
    faculty = (request.GET.get('faculty') or '').strip()
    if faculty and faculty not in srs.get_faculties():
        faculty = ''
    return JsonResponse({'study_programs': srs.get_study_programs(faculty or None)})


def student_intake_template(request):
    """Download the canonical empty workbook used by the importer."""
    require_student_intake_upload(request.user)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Template Upload'
    headers = list(importer.REQUIRED_HEADERS)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='C8102E')
        cell.alignment = Alignment(horizontal='center')
    sheet.freeze_panes = 'A2'
    widths = [16, 38, 58, 12, 16, 12, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    guide = workbook.create_sheet('Petunjuk')
    guide_rows = [
        ['UPLOAD DATA REGISTRASI, KUOTA & BPP'],
        ['Isi data pada sheet "Template Upload". Jangan mengubah nama header.'],
        ['Kode Prodi, Fakultas / Kampus, Program Studi, dan Tahun wajib diisi.'],
        ['Tahun harus berupa angka 4 digit. Tarif, Kuota, Registrasi boleh kosong.'],
        ['Tarif kosong disimpan sebagai NULL; Kuota/Registrasi kosong menjadi 0.'],
        ['Satu Kode Prodi hanya boleh muncul sekali untuk satu Tahun.'],
        ['Upload melakukan UPSERT: baris baru dibuat, baris yang ada diperbarui.'],
        ['Tidak ada penghapusan data lama dari proses upload.'],
        ['Hapus atau ganti contoh sebelum upload bila Anda menambahkan contoh sendiri.'],
    ]
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions['A'].width = 110
    guide['A1'].font = Font(bold=True, size=14, color='C8102E')
    for row in guide.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical='top')

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        'attachment; filename="Template_Upload_Registrasi_Kuota_BPP.xlsx"'
    )
    return response


def _preview_payload(request):
    preview = request.session.get(PREVIEW_SESSION_KEY)
    if not preview or preview.get('expires_at', 0) < time.time():
        request.session.pop(PREVIEW_SESSION_KEY, None)
        return None
    return preview


def student_intake_preview(request):
    """Validate an uploaded workbook and return a preview; never writes trend rows."""
    require_student_intake_upload(request.user)
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'message': 'Gunakan POST.'}, status=405)
    upload = request.FILES.get('file')
    if upload is None:
        return JsonResponse({'ok': False, 'message': 'File Excel wajib dipilih.'}, status=400)
    try:
        rows, summary = importer.preview_upload(upload)
        file_name = importer.safe_file_name(upload.name)
        log = StudentIntakeImportLog.objects.create(
            file_name=file_name,
            uploaded_by=request.user,
            total_rows=summary['total'],
            new_rows=summary['new'],
            updated_rows=summary['update'],
            error_rows=summary['error'],
            status='PREVIEWED',
        )
    except importer.ImportValidationError as exc:
        return JsonResponse({'ok': False, 'message': str(exc)}, status=400)
    except (OperationalError, ProgrammingError):
        return JsonResponse({
            'ok': False,
            'message': 'Database belum siap. Jalankan migrasi sebelum import data.',
        }, status=503)

    token = uuid4().hex
    request.session[PREVIEW_SESSION_KEY] = {
        'token': token,
        'log_id': log.pk,
        'file_name': file_name,
        'rows': rows,
        'summary': summary,
        'expires_at': time.time() + importer.PREVIEW_TTL_SECONDS,
    }
    request.session.modified = True
    return JsonResponse({
        'ok': True,
        'token': token,
        'file_name': file_name,
        'summary': summary,
        'rows': importer.preview_rows_for_response(rows),
        'preview_limit': importer.MAX_PREVIEW_ROWS,
        'can_confirm': summary['error'] == 0,
    })


def student_intake_confirm(request):
    """Confirm the session-owned, already-previewed workbook in one transaction."""
    require_student_intake_upload(request.user)
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'message': 'Gunakan POST.'}, status=405)
    preview = _preview_payload(request)
    if not preview:
        return JsonResponse({'ok': False, 'message': 'Preview sudah kedaluwarsa.'}, status=400)
    token = request.POST.get('token') or request.headers.get('X-Preview-Token')
    if token != preview.get('token'):
        return JsonResponse({'ok': False, 'message': 'Token preview tidak valid.'}, status=400)
    if preview['summary']['error']:
        return JsonResponse({
            'ok': False,
            'message': 'Perbaiki semua baris ERROR sebelum import.',
        }, status=400)

    try:
        result = importer.import_rows(
            preview['rows'], user=request.user,
            source_file_name=preview['file_name'],
        )
        log = StudentIntakeImportLog.objects.get(pk=preview['log_id'])
        log.status = 'SUCCESS'
        log.new_rows = result['new']
        log.updated_rows = result['updated']
        log.error_rows = 0
        log.message = 'Import berhasil.'
        log.save(update_fields=['status', 'new_rows', 'updated_rows', 'error_rows', 'message'])
    except importer.ImportValidationError as exc:
        return JsonResponse({'ok': False, 'message': str(exc)}, status=400)
    except (OperationalError, ProgrammingError):
        return JsonResponse({
            'ok': False,
            'message': 'Database belum siap. Jalankan migrasi sebelum import data.',
        }, status=503)
    except Exception:
        try:
            log = StudentIntakeImportLog.objects.get(pk=preview['log_id'])
            log.status = 'FAILED'
            log.message = 'Import gagal dan seluruh perubahan dibatalkan.'
            log.save(update_fields=['status', 'message'])
        except Exception:
            pass
        return JsonResponse({
            'ok': False,
            'message': 'Import gagal. Tidak ada perubahan yang disimpan.',
        }, status=500)

    request.session.pop(PREVIEW_SESSION_KEY, None)
    request.session.modified = True
    return JsonResponse({'ok': True, **result})