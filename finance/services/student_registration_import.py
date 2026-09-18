"""Safe preview/validation/upsert pipeline for student intake workbooks.

This module runs only from the upload endpoint. It validates the workbook in
memory, returns a JSON-safe preview, and writes rows only after confirmation.
The confirm step is one database transaction.
"""
from collections import Counter
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile, is_zipfile

from django.db import transaction

from finance.models import StudentIntakeTrend
from finance.services.student_registration_service import clear_records_cache

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - declared in requirements.txt
    load_workbook = None

REQUIRED_HEADERS = (
    'Kode Prodi',
    'Fakultas / Kampus',
    'Program Studi',
    'Tahun',
    'Tarif',
    'Kuota',
    'Registrasi',
)
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PREVIEW_ROWS = 100
PREVIEW_TTL_SECONDS = 30 * 60


class ImportValidationError(Exception):
    """A user-correctable workbook error."""


def _clean(value):
    return '' if value is None else str(value).strip()


def _number(value, field, row_number, *, integer=False, nullable=True):
    if value is None or (isinstance(value, str) and not value.strip()):
        if nullable:
            return None
        raise ImportValidationError(f'Baris {row_number}: {field} wajib diisi.')

    text = str(value).strip()
    if isinstance(value, str):
        if ',' in text and '.' in text:
            text = text.replace('.', '').replace(',', '.')
        elif text.count('.') > 1 and ',' not in text:
            text = text.replace('.', '')
        else:
            text = text.replace(',', '.')
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        raise ImportValidationError(
            f'Baris {row_number}: {field} harus berupa angka.') from None
    if not number.is_finite() or number < 0:
        raise ImportValidationError(
            f'Baris {row_number}: {field} harus berupa angka >= 0.')
    if integer and number != number.to_integral_value():
        raise ImportValidationError(
            f'Baris {row_number}: {field} harus berupa bilangan bulat.')
    return int(number) if integer else number


def _year(value, row_number):
    year = _number(value, 'Tahun', row_number, integer=True, nullable=False)
    if year < 1000 or year > 9999:
        raise ImportValidationError(
            f'Baris {row_number}: Tahun harus berupa 4 digit.')
    return year


def _header_map(values):
    headers = [_clean(v) for v in values]
    duplicates = [h for h, count in Counter(headers).items() if h and count > 1]
    if duplicates:
        raise ImportValidationError(
            'Header duplikat tidak diperbolehkan: ' + ', '.join(duplicates))
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ImportValidationError(
            'Kolom wajib tidak ditemukan: ' + ', '.join(missing))
    return {name: headers.index(name) for name in REQUIRED_HEADERS}


def _workbook_bytes(upload):
    if not upload.name.lower().endswith('.xlsx'):
        raise ImportValidationError('File harus berformat .xlsx.')
    if upload.size > MAX_UPLOAD_BYTES:
        raise ImportValidationError('Ukuran file maksimal 10 MB.')
    raw = upload.read()
    upload.seek(0)
    if not raw or not is_zipfile(BytesIO(raw)):
        raise ImportValidationError('File bukan workbook .xlsx yang valid.')
    return raw


def _parse_rows(upload):
    if load_workbook is None:
        raise ImportValidationError('Komponen pembaca Excel belum tersedia di server.')
    raw = _workbook_bytes(upload)
    try:
        workbook = load_workbook(
            BytesIO(raw), read_only=True, data_only=False,
            keep_links=False, keep_vba=False,
        )
    except (BadZipFile, OSError, ValueError) as exc:
        raise ImportValidationError(f'Workbook tidak dapat dibaca: {exc}') from None

    try:
        if 'Template Upload' not in workbook.sheetnames:
            raise ImportValidationError(
                'Sheet wajib "Template Upload" tidak ditemukan.')
        sheet = workbook['Template Upload']
        try:
            header = next(sheet.iter_rows(values_only=False))
        except StopIteration:
            raise ImportValidationError('Sheet "Template Upload" kosong.') from None
        positions = _header_map([cell.value for cell in header])
        parsed = []
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=2, values_only=False), start=2
        ):
            values = [cell.value for cell in cells]
            if not any(_clean(v) for v in values):
                continue
            errors = []
            for field_name, field_index in positions.items():
                cell = cells[field_index] if field_index < len(cells) else None
                value = None if cell is None else cell.value
                if cell is not None and (
                    cell.data_type == 'f'
                    or (isinstance(value, str) and value.lstrip().startswith('='))
                ):
                    errors.append(
                        f'Baris {row_number}: {field_name} tidak boleh berisi formula.'
                    )
            code = _clean(values[positions['Kode Prodi']])
            faculty = _clean(values[positions['Fakultas / Kampus']])
            program = _clean(values[positions['Program Studi']])
            year = None
            tariff = None
            quota = 0
            registration = 0
            try:
                if not code:
                    errors.append('Kode Prodi wajib diisi.')
                if not faculty:
                    errors.append('Fakultas / Kampus wajib diisi.')
                if not program:
                    errors.append('Program Studi wajib diisi.')
                year = _year(values[positions['Tahun']], row_number)
                tariff = _number(values[positions['Tarif']], 'Tarif', row_number)
                quota = _number(
                    values[positions['Kuota']], 'Kuota', row_number, integer=True
                ) or 0
                registration = _number(
                    values[positions['Registrasi']], 'Registrasi', row_number,
                    integer=True,
                ) or 0
            except ImportValidationError as exc:
                errors.append(str(exc).split(': ', 1)[-1])
            parsed.append({
                'row_number': row_number,
                'study_program_code': code,
                'faculty_name': faculty,
                'study_program_name': program,
                'year': year,
                'tariff': None if tariff is None else str(tariff),
                'quota': quota,
                'registration': registration,
                'errors': errors,
            })
        return parsed
    finally:
        workbook.close()


def _key(row):
    return (row['study_program_code'], row['year'])


def _annotate(rows):
    counts = Counter(_key(row) for row in rows if row['study_program_code'] and row['year'])
    duplicate_keys = {key for key, count in counts.items() if count > 1}
    keys = {_key(row) for row in rows if not row['errors'] and _key(row)[1]}
    existing = set(StudentIntakeTrend.objects.filter(
        study_program_code__in={k[0] for k in keys},
        year__in={k[1] for k in keys},
    ).values_list('study_program_code', 'year')) if keys else set()

    summary = {'total': len(rows), 'valid': 0, 'new': 0, 'update': 0, 'error': 0}
    for row in rows:
        key = _key(row)
        if key in duplicate_keys:
            row['errors'].append('Duplikasi Kode Prodi + Tahun pada file upload.')
        if row['errors']:
            row['status'] = 'ERROR'
            row['validation'] = ' '.join(row['errors'])
            summary['error'] += 1
        else:
            row['status'] = 'UPDATE' if key in existing else 'NEW'
            row['validation'] = ''
            summary['valid'] += 1
            summary['update' if row['status'] == 'UPDATE' else 'new'] += 1
    return rows, summary


def preview_upload(upload):
    """Parse, validate and classify a workbook without writing trend rows."""
    rows = _parse_rows(upload)
    return _annotate(rows)


def import_rows(rows, *, user, source_file_name):
    """All-or-nothing upsert for a validated preview payload."""
    if any(row.get('errors') for row in rows):
        raise ImportValidationError('File masih memiliki baris ERROR.')
    created = updated = 0
    with transaction.atomic():
        for row in rows:
            defaults = {
                'faculty_name': row['faculty_name'],
                'study_program_name': row['study_program_name'],
                'tariff': row['tariff'],
                'quota': row['quota'],
                'registration': row['registration'],
                'source_file_name': source_file_name,
                'uploaded_by': user,
            }
            _, was_created = StudentIntakeTrend.objects.update_or_create(
                study_program_code=row['study_program_code'],
                year=row['year'], defaults=defaults,
            )
            if was_created:
                created += 1
            else:
                updated += 1
    clear_records_cache()
    return {'new': created, 'updated': updated, 'total': created + updated}


def preview_rows_for_response(rows):
    """Limit large previews while preserving the all-file summary counts."""
    return rows[:MAX_PREVIEW_ROWS]


def safe_file_name(name):
    return PurePath(name or 'upload.xlsx').name[:255]
